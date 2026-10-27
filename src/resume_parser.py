from pathlib import Path
import re

from pypdf import PdfReader
from docx import Document


# ============================================================
# SECTION DEFINITIONS
# ============================================================

SECTION_ALIASES = {
    "summary": {
        "summary",
        "profile",
        "professional summary",
        "professional profile",
        "career objective",
        "objective",
        "about me",
    },

    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "internship",
        "internships",
        "work history",
    },

    "projects": {
        "projects",
        "academic projects",
        "personal projects",
        "major projects",
        "project experience",
    },

    "skills": {
        "skills",
        "technical skills",
        "core skills",
        "key skills",
        "technical expertise",
        "competencies",
        "technologies",
    },

    "soft_skills": {
        "soft skills",
        "soft skill",
        "personal skills",
        "interpersonal skills",
    },

    "education": {
        "education",
        "academic background",
        "educational background",
        "qualifications",
        "academic qualifications",
    },

    "certifications": {
        "certifications",
        "certificates",
        "licenses",
        "professional certifications",
    },

    "achievements": {
        "achievements",
        "awards",
        "honors",
        "accomplishments",
    },

    "extracurricular": {
        "extracurricular",
        "extra curricular",
        "extra-curricular",
        "extra-curricular activities",
        "extracurricular activities",
    },

    "languages": {
        "languages",
        "language proficiency",
    },

    "interests": {
        "interests",
        "hobbies",
        "personal interests",
    },
}


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_extracted_text(text):
    """Clean and normalize extracted resume text."""

    if not text:
        return ""

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Normalize non-breaking spaces and tabs
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)

    # Normalize excessive blank lines
    text = re.sub(r"\n[ \t]*\n+", "\n\n", text)

    # Remove surrounding whitespace from each line
    lines = [line.strip() for line in text.split("\n")]

    # Remove empty lines
    lines = [line for line in lines if line]

    return "\n".join(lines).strip()


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_text_from_pdf(file_path):
    """Extract text from a PDF resume."""

    reader = PdfReader(file_path)

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text()

            if page_text:
                pages.append(page_text)

        except Exception as e:
            print(
                f"Warning: Could not extract page "
                f"{page_number}: {e}"
            )

    extracted_text = "\n".join(pages)

    return clean_extracted_text(extracted_text)


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(file_path):
    """Extract text from paragraphs and tables in a DOCX resume."""

    document = Document(file_path)

    text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text.append(paragraph_text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text.append(" | ".join(row_text))

    extracted_text = "\n".join(text)

    return clean_extracted_text(extracted_text)


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(file_path):
    """
    Extract readable text from PDF or DOCX.

    This function intentionally continues returning a string
    so the existing ResumeIQ pipeline remains compatible.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if file_path.stat().st_size == 0:
        raise ValueError(
            "The uploaded resume file is empty."
        )

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        text = extract_text_from_pdf(file_path)

    elif extension == ".docx":
        text = extract_text_from_docx(file_path)

    else:
        raise ValueError(
            "Unsupported file type. Please use PDF or DOCX."
        )

    if not text.strip():
        raise ValueError(
            "Could not extract readable text from this resume. "
            "The PDF may be scanned/image-based or the document "
            "may contain unsupported formatting."
        )

    return text


# ============================================================
# SECTION NORMALIZATION
# ============================================================

def normalize_heading(text):
    """Normalize a possible resume section heading."""

    text = text.strip().lower()

    # Remove common decorative characters
    text = re.sub(
        r"^[\s•·▪◦*\-–—:|]+",
        "",
        text
    )

    text = re.sub(
        r"[\s•·▪◦*\-–—:|]+$",
        "",
        text
    )

    # Normalize internal whitespace
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def detect_section_heading(line):
    """
    Detect whether a line represents a known resume section.

    Returns the canonical section name or None.
    """

    normalized = normalize_heading(line)

    if not normalized:
        return None

    # Direct alias match
    for section_name, aliases in SECTION_ALIASES.items():
        if normalized in aliases:
            return section_name

    return None


# ============================================================
# STRUCTURED RESUME SECTIONS
# ============================================================

def parse_resume_sections(resume_text):
    """
    Split resume text into meaningful sections.

    Unknown text before the first recognized heading is stored
    under 'header'.
    """

    if not resume_text:
        return {}

    lines = [
        line.strip()
        for line in resume_text.splitlines()
        if line.strip()
    ]

    sections = {
        "header": []
    }

    current_section = "header"

    for line in lines:

        detected_section = detect_section_heading(line)

        if detected_section:
            current_section = detected_section

            if current_section not in sections:
                sections[current_section] = []

            continue

        sections.setdefault(current_section, [])
        sections[current_section].append(line)

    # Convert lists to clean strings
    cleaned_sections = {}

    for section_name, content in sections.items():

        section_text = "\n".join(content).strip()

        if section_text:
            cleaned_sections[section_name] = section_text

    return cleaned_sections


# ============================================================
# STRUCTURED RESUME PROFILE
# ============================================================

def parse_resume_structure(resume_text):
    """
    Build a structured representation of a resume.

    This does not replace the existing screening engine.
    It provides structured evidence that can be used by
    future ResumeIQ intelligence layers.
    """

    if not resume_text or not resume_text.strip():
        raise ValueError(
            "Cannot structure an empty resume."
        )

    sections = parse_resume_sections(resume_text)

    return {
        "raw_text": resume_text,
        "sections": sections,

        "summary": sections.get("summary", ""),

        "experience": sections.get("experience", ""),

        "projects": sections.get("projects", ""),

        "skills": sections.get("skills", ""),

        "soft_skills": sections.get("soft_skills", ""),

        "education": sections.get("education", ""),

        "certifications": sections.get("certifications", ""),

        "achievements": sections.get("achievements", ""),

        "extracurricular": sections.get(
            "extracurricular",
            ""
        ),

        "languages": sections.get(
            "languages",
            ""
        ),

        "interests": sections.get(
            "interests",
            ""
        ),
    }


# ============================================================
# SIMPLE PROJECT BLOCK EXTRACTION
# ============================================================

def extract_project_blocks(resume_text):
    """
    Extract individual project blocks from the Projects section.

    This intentionally keeps the original project wording.
    Skill identification will continue to use the existing
    ResumeIQ taxonomy rather than introducing a second taxonomy.
    """

    structure = parse_resume_structure(resume_text)

    projects_text = structure.get("projects", "")

    if not projects_text:
        return []

    lines = [
        line.strip()
        for line in projects_text.splitlines()
        if line.strip()
    ]

    projects = []
    current_project = []

    for line in lines:

        # Common project-start pattern:
        # • Project Name
        # - Project Name
        # ▪ Project Name
        #
        # Only start a new block when a bullet begins a line.
        is_new_project = bool(
            re.match(
                r"^[•▪◦*\-–—]\s*",
                line
            )
        )

        if is_new_project and current_project:
            projects.append(
                "\n".join(current_project)
            )
            current_project = []

        current_project.append(line)

    if current_project:
        projects.append(
            "\n".join(current_project)
        )

    return projects